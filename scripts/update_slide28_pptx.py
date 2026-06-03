"""Drop geostrophic-adjustment film-strip into slide 28 of the legoESM seminar deck.

Replaces the 3 placeholder picture boxes with a single wide film-strip showing
SSH evolution at log-spaced times on lat-lon vs MPAS.
"""

from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

DECK = Path("legoESM_seminar (2).pptx")
FILMSTRIP = Path(
    "results/ocean/slide28_validation/geostrophic_adjustment_logtime/"
    "slide28_logtime_filmstrip_eta_u.png"
)

BULLETS = [
    "▸  Geostrophic adjustment from a meridional temperature front — classical Rossby-adjustment validation.",
    "▸  Matched resolution (lat-lon 36×72 ≈ MPAS ico4 ~445 km), matched A_h = 10⁴ m²/s, identical Wright-1997 EOS, identical |lat|>80° land mask.",
    "▸  η panels (top): rest → inertia-gravity ringing (peak ~0.25 m at 6–12 h) → zonally-banded geostrophic balance (~0.06 m by t=10 d).",
    "▸  u panels (bottom): alternating eastward / westward zonal jets emerge by day 4 — the geostrophic flow f·u = −g·∂η/∂y associated with the SSH gradient.",
    "▸  Both grids reproduce the same adjustment trajectory and final balance; visual differences are grid imprint (5° lat-lon stripes vs Voronoi mesh quantization), not numerical dissipation.",
]


def main() -> None:
    backup = DECK.with_suffix(".bak.pptx")
    if not backup.exists():
        shutil.copy2(DECK, backup)
        print(f"backed up to {backup}")

    p = Presentation(str(DECK))
    slide = p.slides[27]  # slide 28, 0-indexed

    # On the pristine slide the relevant shape indices are:
    #   7  Rect 9  (left  picture box, label "MPAS")
    #   8  TextBox "MPAS"
    #   9  Rect 12 (middle picture box, label "FV · lat-lon")
    #  10  TextBox "FV · lat-lon"
    #  11  Rect 15 (right picture box, label "FV · cubed-sphere")
    #  12  TextBox "FV · cubed-sphere"
    #  13  TextBox (bullet body)
    shapes = list(slide.shapes)
    rect_boxes = [shapes[i] for i in (7, 9, 11)]
    panel_labels = [shapes[i] for i in (8, 10, 12)]
    bullet_box = shapes[13]

    # Combined bounding box for the 3 picture boxes — that's where the
    # film-strip goes.
    lefts = [b.left for b in rect_boxes]
    tops = [b.top for b in rect_boxes]
    rights = [b.left + b.width for b in rect_boxes]
    bottoms = [b.top + b.height for b in rect_boxes]
    combined_left = min(lefts)
    combined_top = min(tops)
    combined_width = max(rights) - combined_left
    combined_height = max(bottoms) - combined_top

    # Drop the placeholder rectangles and panel labels — film-strip carries
    # its own grid labels.
    for shape in rect_boxes + panel_labels:
        sp = shape._element
        sp.getparent().remove(sp)

    # Insert the film-strip into the combined box; let pptx preserve aspect
    # by setting width and letting height auto. Then re-position to top of
    # the box so we keep room for the (existing) bullet area below.
    pic = slide.shapes.add_picture(
        str(FILMSTRIP),
        combined_left, combined_top,
        width=combined_width,
    )
    # If the auto-height exceeds the slot, fit by height instead.
    if pic.height > combined_height:
        sp = pic._element
        sp.getparent().remove(sp)
        pic = slide.shapes.add_picture(
            str(FILMSTRIP),
            combined_left, combined_top,
            height=combined_height,
        )
        # Re-centre horizontally inside the slot.
        pic.left = combined_left + (combined_width - pic.width) // 2
    print(f"  inserted film-strip ({pic.width}x{pic.height} EMU)")

    # Bullets — preserve formatting of the first paragraph run, rewrite text.
    from copy import deepcopy
    a_ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    bullet_tf = bullet_box.text_frame
    first_para = bullet_tf.paragraphs[0]
    template_run = first_para.runs[0] if first_para.runs else None
    txBody = bullet_tf._txBody
    for child in list(txBody.findall(f"{a_ns}p")):
        txBody.remove(child)
    for bullet in BULLETS:
        if template_run is not None:
            new_p = deepcopy(first_para._p)
            for r in list(new_p.findall(f"{a_ns}r")):
                new_p.remove(r)
            new_r = deepcopy(template_run._r)
            t = new_r.find(f"{a_ns}t")
            if t is None:
                from lxml import etree
                t = etree.SubElement(new_r, f"{a_ns}t")
            t.text = bullet
            new_p.append(new_r)
            txBody.append(new_p)
        else:
            bullet_tf.add_paragraph().text = bullet
    print(f"  bullets: replaced with {len(BULLETS)} new lines")

    p.save(str(DECK))
    print(f"saved {DECK}")


if __name__ == "__main__":
    main()
